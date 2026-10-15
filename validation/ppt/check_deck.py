"""File-level and negative-path checks; not PowerPoint UI or visual acceptance."""
from pathlib import Path
import contextlib
import importlib.util
import io
import json
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pptx import Presentation

deck, validator, output = map(Path, sys.argv[1:4])
spec = importlib.util.spec_from_file_location('slide_validator', validator)
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)
checks = []
def check(name, condition):
    checks.append({'name':name, 'passed':bool(condition)})
    if not condition:
        raise AssertionError(name)
def validate(p):
    with contextlib.redirect_stdout(io.StringIO()):
        return m.PresentationValidator(str(p)).validate()

prs = Presentation(deck)
check('exactly_three_slides',len(prs.slides)==3)
check('wide_aspect_ratio',abs(prs.slide_width/prs.slide_height-16/9)<1e-5)
texts = ['\n'.join(s.text for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
check('chinese_text_in_native_shapes',all(any('\u4e00' <= c <= '\u9fff' for c in t) for t in texts))
check('synthetic_disclosure_every_slide',all('合成数据' in t for t in texts))
check('no_unsupported_significance_claim','不能据此判断显著性' in texts[1])
check('speaker_notes_present',all(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text for slide in prs.slides))
charts = [s.chart for slide in prs.slides for s in slide.shapes if s.has_chart]
check('one_native_chart',len(charts)==1)
check('chart_numeric_values_preserved',[list(s.values) for s in charts[0].series]==[[20,24,30],[22,30,39]])
check('chart_categories_preserved',[c.label for c in charts[0].plots[0].categories]==['第1周','第2周','第3周'])
check('all_shapes_within_slide',all(s.left>=0 and s.top>=0 and s.left+s.width<=prs.slide_width and s.top+s.height<=prs.slide_height for slide in prs.slides for s in slide.shapes))
with zipfile.ZipFile(deck) as z:
    check('embedded_editable_workbook',any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in z.namelist()))
    check('no_full_slide_raster_fallback',not any(n.startswith('ppt/media/') and not n.endswith('/') for n in z.namelist()))
    check('pptx_zip_crc',z.testzip() is None)
check('valid_deck_accepted',validate(deck)['valid'])
with tempfile.TemporaryDirectory(dir=output.parent) as t:
    t=Path(t)
    e=t/'empty.pptx'; Presentation().save(e)
    check('empty_deck_rejected',not validate(e)['valid'])
    bad=t/'broken.pptx'; bad.write_bytes(b'not a presentation')
    check('corrupt_pptx_rejected',not validate(bad)['valid'])
    check('missing_pptx_rejected',not validate(t/'absent.pptx')['valid'])
    unknown=t/'file.txt'; unknown.write_text('text')
    check('unsupported_extension_rejected',not validate(unknown)['valid'])
    m.HAS_PPTX=False
    check('missing_parser_rejected',not validate(deck)['valid'])
    m.HAS_PPTX=True
result={'scope':'synthetic Chinese editable PPTX; file-level checks only', 'passed':len(checks), 'checks':checks, 'not_tested':['PowerPoint application editing and repair prompts','screen reader reading order','paid image generation','Beamer','real research-data accuracy']}
output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'{len(checks)} checks passed')
